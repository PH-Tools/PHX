# -*- Python Version: 3.10 -*-

"""The closed-file Excel backend: openpyxl, read-only, cached values only.

Implements 'PHX.xl.xl_readable.XLReadable' so that readers written against the
Protocol run against either an open workbook (xlwings 'XLConnection') or a
closed file on disk (this class). Opening is 'read_only=True, data_only=True':
values are the ones Excel last cached — see 'PHX.from_PHPP.staleness' for how
the reader decides whether to trust them.
"""

from __future__ import annotations

import pathlib
from types import TracebackType
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

from PHX.xl import xl_data


class NoSuchSheetError(KeyError):
    """Raised when a worksheet name is not in the workbook (case-insensitive)."""

    def __init__(self, _sheet_name: str, _available: set[str]):
        self.msg = f"Worksheet '{_sheet_name}' not found. Workbook has: {sorted(_available)}"
        super().__init__(self.msg)


class OpenpyxlWorkbook:
    """A closed .xlsx/.xlsm opened read-only through openpyxl, exposing the 'XLReadable' surface.

    Attributes:
        path (pathlib.Path): The workbook on disk. Never written to.
        data_only (bool): True (default) reads cached values; False reads formulas.
    """

    def __init__(self, _path: pathlib.Path | str, data_only: bool = True):
        self.path = pathlib.Path(_path)
        self.data_only = data_only
        self._wb = openpyxl.load_workbook(self.path, read_only=True, data_only=data_only, keep_links=False)
        self._sheet_by_upper: dict[str, str] = {name.upper(): name for name in self._wb.sheetnames}

    # -- context manager --------------------------------------------------

    def __enter__(self) -> OpenpyxlWorkbook:
        return self

    def __exit__(
        self, exc_type: type[BaseException] | None, exc: BaseException | None, tb: TracebackType | None
    ) -> None:
        self.close()

    def close(self) -> None:
        """Release the underlying zip handle."""
        self._wb.close()

    # -- metadata ---------------------------------------------------------

    @property
    def properties(self) -> Any:
        """The openpyxl 'DocumentProperties' (creator, lastModifiedBy, modified, ...)."""
        return self._wb.properties

    @property
    def sheetnames(self) -> list[str]:
        """Worksheet names in workbook order, as written."""
        return list(self._wb.sheetnames)

    @property
    def worksheet_names(self) -> set[str]:
        """Set of the workbook's worksheet names, upper-cased."""
        return set(self._sheet_by_upper)

    # -- reads ------------------------------------------------------------

    def _sheet(self, _sheet_name: str) -> Any:
        key = str(_sheet_name).upper()
        if key not in self._sheet_by_upper:
            raise NoSuchSheetError(_sheet_name, self.worksheet_names)
        return self._wb[self._sheet_by_upper[key]]

    def _read_block(self, _sheet_name: str, _range: str) -> list[list[xl_data.xl_range_single_value]]:
        min_col, min_row, max_col, max_row = range_boundaries(_range)
        ws = self._sheet(_sheet_name)
        rows = ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True)
        block = [list(r) for r in rows]
        # -- read-only sheets stop at their last populated row; pad so the
        # -- caller always gets the full requested shape.
        n_rows = (max_row or min_row) - min_row + 1
        n_cols = (max_col or min_col) - min_col + 1
        while len(block) < n_rows:
            block.append([None] * n_cols)
        return [row + [None] * (n_cols - len(row)) for row in block]

    def get_data(self, _sheet_name: str, _range: str) -> xl_data.xl_writable:
        """Return a value (single cell) or nested lists (range) from the workbook.

        Arguments:
        ----------
            * _sheet_name: (str) The worksheet name (case-insensitive).
            * _range: (str) A cell ("A1") or range ("A1:B4") address.

        Returns:
        --------
            * (xl_writable): scalar for one cell; a flat list for a single row
                or column; a list of lists for a 2-D block (xlwings shapes).
        """
        block = self._read_block(_sheet_name, _range)
        if len(block) == 1 and len(block[0]) == 1:
            return block[0][0]
        if len(block) == 1:
            return block[0]
        if all(len(row) == 1 for row in block):
            return [row[0] for row in block]
        return block  # type: ignore[return-value]

    def get_single_data_item(self, _sheet_name: str, _range: str) -> xl_data.xl_range_single_value:
        """Return the value of a single cell.

        Arguments:
        ----------
            * _sheet_name: (str) The worksheet name (case-insensitive).
            * _range: (str) A single-cell address ("A1").

        Returns:
        --------
            * (xl_range_single_value): The cell's cached value, or None if empty.
        """
        if ":" in _range:
            raise ValueError(f"'get_single_data_item()' takes one cell, got range '{_range}'.")
        return self._read_block(_sheet_name, _range)[0][0]

    def get_single_column_data(
        self,
        _sheet_name: str,
        _col: str,
        _row_start: int | None = None,
        _row_end: int | None = None,
    ) -> list[xl_data.xl_range_single_value]:
        """Return the values of one column between two rows (inclusive).

        Arguments:
        ----------
            * _sheet_name: (str) The worksheet name (case-insensitive).
            * _col: (str) The column letter.
            * _row_start: (int | None) First row (default 1).
            * _row_end: (int | None) Last row (default: the sheet's last row).

        Returns:
        --------
            * (list[xl_range_single_value]): One entry per row, None for empty cells.
        """
        row_start = _row_start or 1
        row_end = _row_end or self._sheet(_sheet_name).max_row or row_start
        block = self._read_block(_sheet_name, f"{_col}{row_start}:{_col}{row_end}")
        return [row[0] for row in block]

    def get_single_row_data(self, _sheet_name: str, _row_number: int) -> list[xl_data.xl_range_single_value]:
        """Return all the values of one row, from column A to the sheet's last used column.

        Arguments:
        ----------
            * _sheet_name: (str) The worksheet name (case-insensitive).
            * _row_number: (int) The row number.

        Returns:
        --------
            * (list[xl_range_single_value]): One entry per column.
        """
        ws = self._sheet(_sheet_name)
        max_col = ws.max_column or column_index_from_string("Z")
        last = get_column_letter(max_col)
        return self._read_block(_sheet_name, f"A{_row_number}:{last}{_row_number}")[0]
