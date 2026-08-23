# -*- Python Version: 3.10 -*-

"""Build small PHPP-shaped workbooks for the public (no client data) tests.

These carry only the sheet names, caption strings and cell positions the reader
depends on — never PHI's workbook itself, which is not redistributed here. Values
are written as plain numbers unless a builder says otherwise, so an openpyxl-written
fixture reads as 'suspect' (non-Excel writer), exactly as the reader should say.
"""

from __future__ import annotations

import pathlib

import openpyxl

from PHX.from_PHPP.results_map import RESULTS_MAP_EN_10_6, ValueKind

PROJECT_VALUES: dict[str, float | str] = {
    "building_type": "1-New building",
    "energy_standard": "10-Passive House",
    "class_pe_method": "10-Classic | PER (renewable)",
    "tfa": 123.4,
    "heating_demand": 14.5,
    "heating_load": 9.8,
    "cooling_demand": "-",
    "overheating_frequency": 3.2,
    "humidity_frequency": 0.0,
    "n50": 0.55,
    "pe_demand": 88.0,
    "per_demand": 41.0,
    "renewable_generation": 12.0,
    "per_demand_without_bioenergy": 45.5,
    "site_energy_demand": 6.7,
    "site_energy_generation": 1.1,
    "cooling_load": 7.9,
}

UNIT_CELLS: dict[str, tuple[str, str, str]] = {  # key -> (sheet, cell, text)
    "tfa": ("Verification", "H35", "m²"),
    "heating_demand": ("Verification", "H36", "kWh/(m²a)"),
    "heating_load": ("Verification", "H37", "W/m²"),
    "cooling_demand": ("Verification", "H39", "kWh/(m²a)"),
    "overheating_frequency": ("Verification", "H40", "%"),
    "humidity_frequency": ("Verification", "H41", "%"),
    "n50": ("Verification", "H43", "1/h"),
    "pe_demand": ("Verification", "H53", "kWh/(m²a)"),
    "per_demand": ("Verification", "H55", "kWh/(m²a)"),
    "renewable_generation": ("Verification", "H56", "kWh/(m²a)"),
    "site_energy_demand": ("PER", "T113", "MWh/a"),
}


def _skeleton(version: str = "10.6", language: str | None = "EN ") -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in ("Brief instructions", "Verification", "Climate", "Areas", "PER", "Cooling load", "Data"):
        wb.create_sheet(name)
    data = wb["Data"]
    data["A3"] = "Data "
    data["A5"] = "PHPP Version"
    data["B5"] = version
    if language is not None:
        data["C5"] = "Language"
        data["D5"] = language
    # -- captions at the mapped label cells
    for spec in RESULTS_MAP_EN_10_6:
        wb[spec.sheet][spec.label_cell] = spec.label_expected
    for _key, (sheet, cell, text) in UNIT_CELLS.items():
        wb[sheet][cell] = text
    wb["Verification"]["J5"] = "Building:"
    wb["Verification"]["K28"] = 20
    wb["Verification"]["N28"] = 25
    wb["Verification"]["T9"] = "1-New building"
    wb["Verification"]["T11"] = "10-Passive House"
    wb["Verification"]["T14"] = "10-Classic | PER (renewable)"
    wb["Climate"]["D12"] = "DE-9999-PHPP-Standard"
    return wb


def write_blank(path: pathlib.Path) -> pathlib.Path:
    """A PHPP-shaped blank template: captions, template defaults, TFA 0, no project content."""
    wb = _skeleton()
    wb["Verification"]["I35"] = 0
    wb["Verification"]["I53"] = 0
    wb["Verification"]["I55"] = 0
    wb.save(path)
    return path


def write_project(
    path: pathlib.Path,
    values: dict[str, float | str] | None = None,
    as_formulas: bool = False,
    move_caption: str | None = None,
    drop_sheet: str | None = None,
) -> pathlib.Path:
    """A PHPP-shaped project workbook.

    Arguments:
    ----------
        * path: where to save.
        * values: per-key values (default PROJECT_VALUES).
        * as_formulas: write the numeric results as formulas (openpyxl leaves them uncached).
        * move_caption: a results key whose caption is replaced by 'Something else' (label mismatch).
        * drop_sheet: a sheet name to remove (ie: 'Cooling load').
    """
    vals = dict(PROJECT_VALUES if values is None else values)
    wb = _skeleton()
    wb["Verification"]["F29"] = 1
    wb["Verification"]["F28"] = 2026
    wb["Climate"]["D12"] = "XX-0001-Testville"
    wb["Areas"]["L8"] = vals.get("tfa", 0)
    for spec in RESULTS_MAP_EN_10_6:
        v = vals.get(spec.key)
        if v is None:
            continue
        if as_formulas and spec.kind is ValueKind.NUMBER and not isinstance(v, str):
            wb[spec.sheet][spec.cell] = f"={v}"
        else:
            wb[spec.sheet][spec.cell] = v
    if move_caption:
        spec = next(s for s in RESULTS_MAP_EN_10_6 if s.key == move_caption)
        wb[spec.sheet][spec.label_cell] = "Something else"
    if drop_sheet:
        wb.remove(wb[drop_sheet])
    wb.save(path)
    return path


def write_version(path: pathlib.Path, version: str, language: str | None) -> pathlib.Path:
    """A PHPP-shaped project with a given 'Data' version row (language None = PHPP 9 layout)."""
    wb = _skeleton(version=version, language=language)
    wb["Verification"]["F29"] = 2
    wb["Verification"]["I35"] = 99.0
    wb.save(path)
    return path


def write_not_a_phpp(path: pathlib.Path) -> pathlib.Path:
    """A plain workbook with none of the PHPP sheets."""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "PHPP Version"  # a decoy string in the wrong place
    wb.active["B1"] = "10.6"
    wb.save(path)
    return path
