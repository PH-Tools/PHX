# -*- Python Version: 3.10 -*-

import pathlib

import openpyxl
import pytest

from PHX.from_PHPP.xl_openpyxl import NoSuchSheetError, OpenpyxlWorkbook
from PHX.xl.xl_app import XLConnection
from PHX.xl.xl_readable import XLReadable


def _book(tmp_path: pathlib.Path) -> pathlib.Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r, row in enumerate([["a", 1, 2.5], ["b", None, "x"], [None, None, None], ["d", 4, 5]], start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)
    p = tmp_path / "b.xlsx"
    wb.save(p)
    return p


def test_openpyxl_backend_satisfies_the_read_protocol(tmp_path: pathlib.Path) -> None:
    with OpenpyxlWorkbook(_book(tmp_path)) as xl:
        assert isinstance(xl, XLReadable)
    # -- and so does the xlwings facade, structurally (no instance needed for the attribute check)
    for name in (
        "get_data",
        "get_single_data_item",
        "get_single_column_data",
        "get_single_row_data",
        "worksheet_names",
    ):
        assert hasattr(XLConnection, name)


def test_openpyxl_backend_read_shapes(tmp_path: pathlib.Path) -> None:
    with OpenpyxlWorkbook(_book(tmp_path)) as xl:
        assert xl.worksheet_names == {"DATA"}
        assert xl.get_single_data_item("data", "B1") == 1  # case-insensitive sheet
        assert xl.get_data("Data", "A1") == "a"
        assert xl.get_data("Data", "A1:C1") == ["a", 1, 2.5]
        assert xl.get_data("Data", "A1:A4") == ["a", "b", None, "d"]
        assert xl.get_data("Data", "A1:B2") == [["a", 1], ["b", None]]
        assert xl.get_single_column_data("Data", "B", 1, 4) == [1, None, None, 4]
        # -- rows past the sheet's last populated row are padded, never dropped
        assert xl.get_single_column_data("Data", "B", 3, 6) == [None, 4, None, None]
        assert xl.get_single_row_data("Data", 4)[:3] == ["d", 4, 5]
        with pytest.raises(ValueError):
            xl.get_single_data_item("Data", "A1:B2")
        with pytest.raises(NoSuchSheetError):
            xl.get_data("Nope", "A1")
